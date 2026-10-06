"""對「測試副本」的實機整合測試：讀取 → 貼上 → 讀回 → 下載檢查字型 → Ctrl+Z 復原。

只在手動設定環境變數時執行，絕不可指向實驗室正式表單：
  INSTRUMENT_BOOKING_LIVE_URL     測試副本網址
  INSTRUMENT_BOOKING_LIVE_RANGES  以分號分隔的空白範圍，例 "202610!U60:U61;202610!U63:U63"
  INSTRUMENT_BOOKING_PROFILE      （選用）Edge 設定檔資料夾，預設 %LOCALAPPDATA%\\InstrumentBooking\\edge-profile
"""
import os
import time
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import range_boundaries

from instrument_booking.browser.downloader import SnapshotDownloader, file_id_from_url
from instrument_booking.browser.session import EdgeSession, default_profile_dir
from instrument_booking.browser.sheets_writer import BrowserSheetWriter
from instrument_booking.core.models import CellFont, CellState
from instrument_booking.core.occupancy import is_empty

URL = os.environ.get("INSTRUMENT_BOOKING_LIVE_URL", "")
RANGES = [r for r in os.environ.get("INSTRUMENT_BOOKING_LIVE_RANGES", "").split(";") if r]
PROFILE = Path(os.environ["INSTRUMENT_BOOKING_PROFILE"]) if os.environ.get("INSTRUMENT_BOOKING_PROFILE") else None

pytestmark = [
    pytest.mark.live,
    pytest.mark.skipif(not URL or not RANGES, reason="未設定實機測試環境變數"),
]

LAB = CellFont(12.0, True, "center")
NAME = "IT-TEST"


class SystemClock:
    def now(self):
        return time.time()


@pytest.fixture(scope="module")
def session():
    s = EdgeSession(URL, PROFILE or default_profile_dir())
    s.start()
    yield s
    s.close()


@pytest.mark.parametrize("ref", RANGES)
def test_read_paste_verify_and_undo(session, ref, tmp_path):
    sheet, a1 = ref.split("!")
    page = session.sheet_page()
    writer = BrowserSheetWriter(page, restart=session.restart_sheet_page, clock=SystemClock(), not_before=0.0)
    before = writer.read_range(sheet, a1)
    assert all(is_empty(c) for c in before), f"{ref} 不是空的，請改用空白範圍"
    pasted = False
    try:
        writer.paste_booking(sheet, a1, "A4C2F4", NAME, [LAB] * len(before))
        pasted = True
        page.wait(1500)
        after = writer.read_range(sheet, a1)
        assert after == [CellState(NAME, "A4C2F4")] + [CellState(None, "A4C2F4")] * (len(before) - 1)

        snapshot = SnapshotDownloader(lambda: session.request, file_id_from_url(URL), tmp_path).download()
        ws = openpyxl.load_workbook(snapshot)[sheet]
        col, row1, _, row2 = range_boundaries(a1)
        for row in range(row1, row2 + 1):
            cell = ws.cell(row, col)
            assert (cell.font.sz, bool(cell.font.b), cell.alignment.horizontal) == (12.0, True, "center"), cell.coordinate
        assert ws.cell(row2 + 1, col).fill.fgColor.rgb != "FFA4C2F4", "貼上超出範圍"
    finally:
        if pasted:
            page.press("Control+Z")
            page.wait(1500)
    restored = writer.read_range(sheet, a1)
    assert all(is_empty(c) for c in restored), f"Ctrl+Z 未完全復原 {ref}，請用版本記錄還原測試副本"
