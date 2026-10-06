"""建立與實驗室預約表相同結構的小型測試活頁簿（結構見 docs/表格結構.md §3、§4）。"""
from __future__ import annotations

from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import PatternFill

DEFAULT_LEGEND = [
    ("H2", "F4CCCC"), ("O2", "B6D7A8"), ("Ar", "A4C2F4"), ("vac", "FFE599"), ("air", "B4A7D6"),
    ("MILA", "999999"), ("Ar/H2", "DD7E6B"), ("CO2/H2", "76A5AF"), ("N2", "D5A6BD"),
]
TIME_LABELS = [f"{h:02d}00~{h + 1:02d}00" for h in range(9, 23)] + ["2300~"]
WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

Layout = tuple[int, int, dict[int, int]]  # (date_row, first_col, {hour: row})


def solid(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor="FF" + rgb)


def new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _write_weeks(ws, *, first_date_row, monday, weeks, widths, blank_row_after_hour) -> dict[date, Layout]:
    layout: dict[date, Layout] = {}
    row = first_date_row
    for w in range(weeks):
        date_row = row
        hour_rows: dict[int, int] = {}
        r = date_row + 1
        for i, label in enumerate(TIME_LABELS):
            hour = 9 + i
            ws.cell(r, 1, label)
            hour_rows[hour] = r
            r += 1
            if blank_row_after_hour == hour:
                r += 1  # 誤插入的列：A 欄空白
        col = 2
        for d in range(7):
            day = monday + timedelta(days=7 * w + d)
            ws.cell(date_row, col, datetime.combine(day, datetime.min.time()))
            ws.merge_cells(start_row=date_row, start_column=col, end_row=date_row, end_column=col + widths[d] - 1)
            layout[day] = (date_row, col, dict(hour_rows))
            col += widths[d]
        row = r
    return layout


def add_tube_sheet(wb, title, monday, *, weeks=1, legend=DEFAULT_LEGEND, widths=(3,) * 7,
                   blank_row_after_hour=None):
    """管型爐表：第 2 列圖例、第 3 列 A/B/C、第 4 列星期、第 5 列起每週 16 列。"""
    ws = wb.create_sheet(title)
    ws["A2"] = "°C"
    for i, (name, rgb) in enumerate(legend):
        cell = ws.cell(2, 2 + i, name)
        if rgb is not None:
            cell.fill = solid(rgb)
    ws["A3"] = "tube"
    col = 2
    for d in range(7):
        for k in range(widths[d]):
            ws.cell(3, col + k, "ABC"[k] if k < 3 else "反應器")
        ws.cell(4, col, WEEKDAYS[d])
        col += widths[d]
    layout = _write_weeks(ws, first_date_row=5, monday=monday, weeks=weeks, widths=widths,
                          blank_row_after_hour=blank_row_after_hour)
    return ws, layout


def add_oven_sheet(wb, title, monday, *, weeks=1, header=True, blank_row_after_hour=None):
    """烘箱表：第 1 列 A/B（ovenA 黃、ovenB 藍）、第 2 列星期與 P2/Q2 圖例、第 3 列起每週 16 列。"""
    ws = wb.create_sheet(title)
    for d in range(7):
        col = 2 + 2 * d
        if header:
            ws.cell(1, col, "A").fill = solid("FFE599")
            ws.cell(1, col + 1, "B").fill = solid("A4C2F4")
        ws.cell(2, col, WEEKDAYS[d])
    ws["P2"] = "ovenA"
    ws["P2"].fill = solid("FFE599")
    ws["Q2"] = "ovenB"
    ws["Q2"].fill = solid("A4C2F4")
    layout = _write_weeks(ws, first_date_row=3, monday=monday, weeks=weeks, widths=(2,) * 7,
                          blank_row_after_hour=blank_row_after_hour)
    return ws, layout
