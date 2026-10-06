import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.colors import Color

from instrument_booking.core.models import CellFont, CellState
from instrument_booking.core.occupancy import cell_font, cell_state, is_empty, normalize_rgb


@pytest.fixture
def ws():
    return Workbook().active


def test_normalize_rgb():
    assert normalize_rgb("FFA4C2F4") == "A4C2F4"
    assert normalize_rgb("a4c2f4") == "A4C2F4"


def test_no_fill_no_value_is_empty(ws):
    s = cell_state(ws["B6"])
    assert s == CellState(None, None)
    assert is_empty(s)


def test_white_rgb_is_empty(ws):
    ws["B6"].fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    assert is_empty(cell_state(ws["B6"]))


def test_theme0_white_is_empty(ws):
    ws["B6"].fill = PatternFill(fill_type="solid", fgColor=Color(theme=0))
    assert is_empty(cell_state(ws["B6"]))


def test_theme0_with_tint_is_occupied(ws):
    ws["B6"].fill = PatternFill(fill_type="solid", fgColor=Color(theme=0, tint=-0.15))
    s = cell_state(ws["B6"])
    assert s.color == "theme:0"
    assert not is_empty(s)


def test_colored_cell_is_occupied_and_normalized(ws):
    ws["B6"].fill = PatternFill(fill_type="solid", fgColor="FFA4C2F4")
    s = cell_state(ws["B6"])
    assert s.color == "A4C2F4"
    assert not is_empty(s)


def test_indexed_color_is_occupied(ws):
    ws["B6"].fill = PatternFill(fill_type="solid", fgColor=Color(indexed=5))
    assert not is_empty(cell_state(ws["B6"]))


@pytest.mark.parametrize("value", ["Zoe", 70.0, 0])
def test_any_value_is_occupied(ws, value):
    ws["B6"] = value
    assert not is_empty(cell_state(ws["B6"]))


def test_empty_string_is_empty():
    assert is_empty(CellState("", None))


def test_merged_inner_cell_is_occupied(ws):
    ws.merge_cells("B6:B7")
    s = cell_state(ws["B7"])
    assert s.color == "merged"
    assert not is_empty(s)


def test_cell_font_reads_lab_style(ws):
    # 實驗室預約格慣例：12pt 粗體置中
    ws["B6"].font = Font(size=12, bold=True)
    ws["B6"].alignment = Alignment(horizontal="center")
    assert cell_font(ws["B6"]) == CellFont(12.0, True, "center")


def test_cell_font_default(ws):
    assert cell_font(ws["B6"]) == CellFont(11.0, False, None)  # openpyxl 新活頁簿預設 11pt


def test_cell_font_merged_inner_cell(ws):
    ws.merge_cells("B6:B7")
    assert cell_font(ws["B7"]) == CellFont(None, False, None)
