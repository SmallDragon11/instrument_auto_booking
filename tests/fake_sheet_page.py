"""模擬 Google 試算表網頁的 SheetPage：名稱方塊跳轉、Ctrl+C 產生與 Google 相同格式的 HTML、Ctrl+V 寫回。"""
from __future__ import annotations

import html as html_lib
import re

from openpyxl.utils import get_column_letter, range_boundaries

from instrument_booking.core.clipboard_html import parse_cells
from instrument_booking.core.models import CellState

_REF = re.compile(r"^'(?P<sheet>[^']+)'!(?P<a1>[A-Z]+\d+(?::[A-Z]+\d+)?)$")

WRAPPER_OPEN = ('<google-sheets-html-origin style="color: rgb(0, 0, 0); font-size: medium;">'
                '<table xmlns="http://www.w3.org/1999/xhtml" cellspacing="0" cellpadding="0" dir="ltr" border="1" '
                'data-sheets-root="1" data-sheets-baot="1" style="table-layout: fixed; font-size: 10pt; '
                'font-family: Arial;"><colgroup><col width="46"></colgroup><tbody>')
WRAPPER_CLOSE = "</tbody></table></google-sheets-html-origin>"
SPAN = ('<span data-sheets-root="1" style="font-size: 12pt; font-family: Arial; font-weight: bold; '
        'color: rgb(255, 0, 0);"></span>')


def _td(s: CellState) -> str:
    rgb = s.color or "FFFFFF"
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    style = (f"border-width: 1px; border-style: solid; overflow: hidden; padding: 0px 3px; "
             f"vertical-align: bottom; background-color: rgb({r}, {g}, {b});")
    if s.value not in (None, ""):
        style += " font-size: 12pt; font-weight: bold; text-align: center;"
    text = html_lib.escape(str(s.value)) if s.value not in (None, "") else ""
    return f'<td style="{style}">{text}</td>'


def google_html_rows(rows: list[list[CellState]]) -> str:
    """產生與 Google 試算表相同結構的複製 HTML；每列可有多格（多欄選取）。"""
    trs = "".join(f'<tr style="height: 21px;">{"".join(_td(s) for s in row)}</tr>' for row in rows)
    return WRAPPER_OPEN + trs + WRAPPER_CLOSE


def google_html(states: list[CellState]) -> str:
    """產生與 Google 試算表相同結構的多格（單欄）複製 HTML（白色以 rgb(255, 255, 255) 表示）。"""
    return google_html_rows([[s] for s in states])


class FakeSheetPage:
    def __init__(self) -> None:
        self.cells: dict[tuple[str, str], CellState] = {}
        self.selection: tuple[str, str] | None = None
        self.clipboard_html: str | None = None
        self.clipboard_text = ""
        self.pasted: list[tuple[str, str, str]] = []  # (sheet, 第一格, html)
        self.log: list[tuple] = []
        self.elapsed = 0.0        # 假時間（秒），只由 wait() 推進
        self.stale_copies = 0     # 接下來幾次 Ctrl+C 只複製到單格 <span>（模擬跳轉尚未生效）
        self.name_box_lag = 0     # 接下來幾次 name_box() 回傳舊位置
        self.active_lag = 0       # 接下來幾次 active_sheet() 回傳「切換中」（模擬工作表尚未切換完成）
        self.fail: dict[str, Exception] = {}  # 操作名稱 → 下一次呼叫時拋出的例外
        # 工作表 → 合併列（例如下一週的日期列，整列橫跨該天三欄）。選取包含這些列的單欄範圍時，
        # Google 會把選取擴張成左右各一欄（實測：選 U20:U21 → 名稱方塊顯示 T20:V21）
        self.merged_rows: dict[str, set[int]] = {}

    def now(self) -> float:
        return self.elapsed

    def _maybe_fail(self, op: str) -> None:
        exc = self.fail.pop(op, None)
        if exc is not None:
            raise exc

    def jump(self, ref: str) -> None:
        self._maybe_fail("jump")
        m = _REF.match(ref)
        assert m, f"名稱方塊格式錯誤：{ref}"
        self.selection = (m["sheet"], self._expand(m["sheet"], m["a1"]))
        self.log.append(("jump", ref))

    def _expand(self, sheet: str, a1: str) -> str:
        c1, r1, c2, r2 = range_boundaries(a1)
        merged = self.merged_rows.get(sheet, set())
        if c1 != c2 or not any(r in merged for r in range(r1, r2 + 1)):
            return a1
        return f"{get_column_letter(max(c1 - 1, 1))}{r1}:{get_column_letter(c2 + 1)}{r2}"

    def name_box(self) -> str:
        if self.name_box_lag:
            self.name_box_lag -= 1
            return "A1"
        return self.selection[1] if self.selection else ""

    def active_sheet(self) -> str:
        if self.active_lag:
            self.active_lag -= 1
            return "（切換中）"
        return self.selection[0] if self.selection else ""

    def press(self, keys: str) -> None:
        self._maybe_fail(keys)
        self.log.append(("press", keys))
        sheet, a1 = self.selection
        if keys == "Control+C":
            c1, r1, c2, r2 = range_boundaries(a1)
            if self.stale_copies:
                self.stale_copies -= 1
                self.clipboard_html = SPAN
            elif (c1, r1) == (c2, r2):
                self.clipboard_html = SPAN  # 真實行為：單格沒有底色資訊
            else:
                self.clipboard_html = google_html_rows([
                    [self.cells.get((sheet, f"{get_column_letter(c)}{r}"), CellState(None, None))
                     for c in range(c1, c2 + 1)]
                    for r in range(r1, r2 + 1)])
        elif keys == "Control+V":
            states = parse_cells(self.clipboard_html)
            col1, row1, _, _ = range_boundaries(a1)
            for i, s in enumerate(states):
                self.cells[(sheet, f"{get_column_letter(col1)}{row1 + i}")] = s
            self.pasted.append((sheet, a1, self.clipboard_html))

    def read_clipboard_html(self) -> str | None:
        self._maybe_fail("read_clipboard_html")
        return self.clipboard_html

    def write_clipboard(self, html: str | None, text: str) -> None:
        self._maybe_fail("write_clipboard")
        self.clipboard_html = html or None
        self.clipboard_text = text

    def wait(self, ms: int) -> None:
        self.elapsed += ms / 1000
