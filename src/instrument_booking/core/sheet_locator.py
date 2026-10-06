"""依日期定位預約表中的儲存格（規則見 docs/表格結構.md §6）。"""
from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from typing import Mapping

from openpyxl.utils import get_column_letter

from instrument_booking.core.models import FIRST_HOUR, LAST_HOUR, BookingRequest

TUBE_SHEET = re.compile(r"^\d{6}$")
OVEN_SHEET = re.compile(r"^\d{1,2}月oven(?:\d{4})?$")
TIME_LABEL = re.compile(r"^(\d{2})00~(?:\d{2}00)?$")
SUB_COUNT = {"tube": 3, "oven": 2}
KIND_NAME = {"tube": "管型爐", "oven": "烘箱"}
ALL_HOURS = tuple(range(FIRST_HOUR, LAST_HOUR))


def sheet_kind(title: str) -> str | None:
    t = title.strip()
    if TUBE_SHEET.match(t):
        return "tube"
    if OVEN_SHEET.match(t):
        return "oven"
    return None


def parse_time_label(value: object) -> int | None:
    if value is None:
        return None
    m = TIME_LABEL.match(str(value).strip())
    if not m:
        return None
    hour = int(m.group(1))
    return hour if FIRST_HOUR <= hour < LAST_HOUR else None


@dataclass(frozen=True)
class DayBlock:
    sheet: str
    date: date
    first_col: int
    width: int
    date_row: int
    hour_rows: Mapping[int, int]

    def row_of(self, hour: int) -> int:
        return self.hour_rows[hour]


@dataclass(frozen=True)
class Target:
    sheet: str
    column: int
    first_row: int
    last_row: int

    @property
    def rows(self) -> range:
        return range(self.first_row, self.last_row + 1)

    def cell_name(self, row: int) -> str:
        return f"{get_column_letter(self.column)}{row}"

    @property
    def first_cell(self) -> str:
        return self.cell_name(self.first_row)

    @property
    def a1(self) -> str:
        return f"{self.first_cell}:{self.cell_name(self.last_row)}"


class LocateError(Exception):
    """無法定位（訊息為給使用者看的繁體中文）。"""


class SheetMissing(LocateError):
    pass


class AmbiguousDate(LocateError):
    pass


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _scan_sheet(ws, kind: str) -> list[DayBlock]:
    width_needed = SUB_COUNT[kind]
    merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
    found: list[tuple[int, int, int, date]] = []  # (row, col, width, date)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            d = _as_date(cell.value)
            if d is None:
                continue
            m = merged.get((cell.row, cell.column))
            if m is None:
                continue
            width = m.max_col - m.min_col + 1
            if width < width_needed:
                continue
            found.append((cell.row, cell.column, width, d))
    date_rows = sorted({r for r, _, _, _ in found})
    blocks = []
    for r, col, width, d in found:
        next_row = next((x for x in date_rows if x > r), ws.max_row + 1)
        hour_rows: dict[int, int] = {}
        for rr in range(r + 1, next_row):
            hour = parse_time_label(ws.cell(rr, 1).value)
            if hour is not None and hour not in hour_rows:
                hour_rows[hour] = rr
        rows_in_order = [hour_rows.get(h) for h in ALL_HOURS]
        if None in rows_in_order or rows_in_order != sorted(rows_in_order):
            continue  # 下方沒有完整時段列 → 不是預約格的日期
        blocks.append(DayBlock(ws.title, d, col, width, r, hour_rows))
    return blocks


class SheetIndex:
    def __init__(self, blocks: dict[tuple[str, date], list[DayBlock]]) -> None:
        self._blocks = blocks

    @classmethod
    def build(cls, wb) -> SheetIndex:
        blocks: dict[tuple[str, date], list[DayBlock]] = defaultdict(list)
        for ws in wb.worksheets:
            kind = sheet_kind(ws.title)
            if kind is None:
                continue
            for block in _scan_sheet(ws, kind):
                blocks[(kind, block.date)].append(block)
        return cls(dict(blocks))

    def all_blocks(self) -> list[DayBlock]:
        return [b for bs in self._blocks.values() for b in bs]

    def day(self, kind: str, d: date) -> DayBlock:
        found = self._blocks.get((kind, d), [])
        if not found:
            raise SheetMissing(f"找不到 {d:%m/%d} 所在的{KIND_NAME[kind]}工作表")
        if len(found) > 1:
            where = "、".join(f"{b.sheet}!{get_column_letter(b.first_col)}{b.date_row}" for b in found)
            raise AmbiguousDate(f"{d:%m/%d} 同時出現在 {where}，無法判斷要寫入哪一處")
        return found[0]

    def locate(self, req: BookingRequest) -> Target:
        block = self.day(req.instrument.kind, req.date)
        return Target(
            sheet=block.sheet,
            column=block.first_col + req.instrument.sub_index,
            first_row=block.row_of(req.start_hour),
            last_row=block.row_of(req.end_hour - 1),
        )
